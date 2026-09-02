#!/usr/bin/env python3
"""
Fast layout-preserving PDF → DOCX conversion (local, no cloud APIs).

Pipeline:
  1. Profile the PDF (text, images, scanned-page detection).
  2. Convert the whole document in one pdf2docx pass (fonts, tables, alignment).
  3. If that fails, reconstruct page-by-page with image fallback.
  4. Apply fidelity cleanup + monospace/code styling.
  5. Validate output against the source profile.
"""

from __future__ import annotations

import json
import sys
import time
from pathlib import Path

from docx_page_builder import build_fidelity_docx
from docx_validator import validate_document
from pdf_analysis import analyze_pdf, DocumentProfile


def log(msg: str) -> None:
    print(msg, file=sys.stderr, flush=True)


def run_ocr_if_needed(pdf_path: Path, likely_scanned: bool) -> Path:
    if not likely_scanned:
        return pdf_path

    try:
        import ocrmypdf
    except ImportError:
        log("ocrmypdf not installed — continuing without OCR (scanned PDF quality may be lower)")
        return pdf_path

    ocr_path = pdf_path.with_name(f"{pdf_path.stem}_ocr.pdf")
    log("running OCR for scanned PDF…")
    try:
        ocrmypdf.ocr(
            str(pdf_path),
            str(ocr_path),
            skip_text=True,
            optimize=0,
            output_type="pdf",
            progress_bar=False,
        )
        return ocr_path
    except Exception as exc:
        log(f"OCR skipped: {exc}")
        if ocr_path.exists():
            ocr_path.unlink(missing_ok=True)
        return pdf_path


def fidelity_cleanup(docx_path: Path) -> dict:
    from docx import Document

    from docx_cleanup import fidelity_docx_cleanup

    doc = Document(str(docx_path))
    stats = fidelity_docx_cleanup(doc)
    doc.save(str(docx_path))
    if any(stats.values()):
        log(f"fidelity cleanup: {stats}")
    return stats


def apply_monospace_styles(docx_path: Path, snippets: list[str], code_lines: list[str]) -> int:
    if not snippets and not code_lines:
        return 0

    from docx import Document
    from docx.oxml.ns import qn

    doc = Document(str(docx_path))
    mono_font = "Consolas"
    styled = 0
    unique = sorted(
        {s for s in snippets + code_lines if s and len(s) >= 2},
        key=len,
        reverse=True,
    )

    def style_paragraph(para) -> None:
        nonlocal styled
        para_text = para.text
        if not para_text.strip():
            return
        for snippet in unique:
            if snippet not in para_text:
                continue
            for run in para.runs:
                if snippet in run.text or run.text in snippet:
                    run.font.name = mono_font
                    r = run._element.get_or_add_rPr()
                    r.rFonts.set(qn("w:ascii"), mono_font)
                    r.rFonts.set(qn("w:hAnsi"), mono_font)
                    r.rFonts.set(qn("w:cs"), mono_font)
                    r.rFonts.set(qn("w:eastAsia"), mono_font)
                    styled += 1

    for para in doc.paragraphs:
        style_paragraph(para)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    style_paragraph(para)

    doc.save(str(docx_path))
    return styled


def convert_single_pass(pdf_path: Path, docx_path: Path, log) -> dict:
    """Fast layout-preserving conversion of the whole PDF in one pdf2docx pass."""
    from pdf2docx import Converter

    log("running pdf2docx single-pass layout reconstruction…")
    started = time.monotonic()
    converter = Converter(str(pdf_path))
    try:
        converter.convert(str(docx_path))
    finally:
        converter.close()

    if not docx_path.exists() or docx_path.stat().st_size < 200:
        raise RuntimeError("pdf2docx produced an empty document")

    elapsed_ms = int((time.monotonic() - started) * 1000)
    log(f"single-pass conversion finished in {elapsed_ms}ms")
    return {"mode": "single-pass", "elapsed_ms": elapsed_ms}


def convert_fast_then_fallback(
    pdf_path: Path,
    docx_path: Path,
    profile: DocumentProfile,
    log,
) -> dict:
    try:
        stats = convert_single_pass(pdf_path, docx_path, log)
        stats.update(
            {
                "page_modes": ["vector"] * profile.page_count,
                "image_fallback_pages": [],
                "vector_pages": profile.page_count,
                "image_pages": 0,
            }
        )
        return stats
    except Exception as exc:
        log(f"single-pass failed ({exc}); falling back to per-page reconstruction")
        if docx_path.exists():
            docx_path.unlink(missing_ok=True)
        return build_fidelity_docx(pdf_path, docx_path, profile, log)


def convert(input_pdf: str, output_docx: str) -> dict:
    src = Path(input_pdf).resolve()
    dst = Path(output_docx).resolve()
    dst.parent.mkdir(parents=True, exist_ok=True)

    if not src.exists():
        raise FileNotFoundError(f"Input PDF not found: {src}")

    profile = analyze_pdf(src)
    agg = profile.aggregate()
    log(
        f"PDF profile: {agg['page_count']} pages, {agg['text_chars']} chars, "
        f"{agg['image_count']} images, {agg['heading_count']} headings, "
        f"{agg['code_block_count']} code blocks"
    )

    work_pdf = run_ocr_if_needed(src, profile.likely_scanned)
    ocr_used = work_pdf != src

    try:
        if work_pdf != src:
            profile = analyze_pdf(work_pdf)

        build_stats = convert_fast_then_fallback(work_pdf, dst, profile, log)
        cleanup_stats = fidelity_cleanup(dst)

        mono_snippets = list(
            set(profile.monospace_snippets + profile.code_lines)
        )
        styled_runs = apply_monospace_styles(dst, mono_snippets, profile.code_lines)

        from docx import Document

        doc = Document(str(dst))
        validation = validate_document(
            profile,
            doc,
            page_modes=build_stats.get("page_modes"),
        )

        if not validation.passed:
            log(f"validation warnings/failures: {validation.failures}")
            if validation.warnings:
                log(f"validation notes: {validation.warnings}")

        table_count = len(doc.tables)
        image_count = sum(
            1 for rel in doc.part.rels.values() if "image" in rel.reltype
        )

        return {
            "ok": True,
            "output": str(dst),
            "page_count": profile.page_count,
            "byte_length": dst.stat().st_size,
            "pdf_image_count": profile.image_count,
            "docx_image_count": image_count,
            "docx_table_count": table_count,
            "monospace_runs_styled": styled_runs,
            "ocr_used": ocr_used,
            "engine": "pdf2docx-fidelity",
            "build": build_stats,
            "cleanup": cleanup_stats,
            "validation": validation.to_dict(),
        }
    finally:
        if ocr_used and work_pdf.exists():
            work_pdf.unlink(missing_ok=True)


def main() -> int:
    if len(sys.argv) != 3:
        print(json.dumps({"ok": False, "error": "Usage: pdf_to_docx.py <input.pdf> <output.docx>"}))
        return 1

    try:
        result = convert(sys.argv[1], sys.argv[2])
        print(json.dumps(result))
        return 0
    except Exception as exc:
        print(json.dumps({"ok": False, "error": str(exc)}))
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
